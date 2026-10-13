"""可信性评估报告导出。

先把汇总结论 + 各模块明细装配成一份结构化 `ReportModel`（内容只写一遍），
再交给两个渲染器：
    - PDF ：PyMuPDF 内置简体中文字体（china-s），跨平台一致，无需系统字体；
    - Word：python-docx。

数字与等级全部来自 `verdict` 层的既有判定，渲染器只负责排版，不做任何推断。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field

from app.core.config import settings
from app.modules.verdict import service
from app.modules.verdict.schema import OverallVerdict

try:  # PyMuPDF >= 1.24 推荐 pymupdf；旧版仅有 fitz
    import pymupdf as fitz
except ImportError:  # pragma: no cover
    import fitz  # type: ignore

FONT = "china-s"  # PyMuPDF 内置简体中文字体
A4_W, A4_H = 595, 842
MARGIN_X, MARGIN_Y = 50, 56

# 等级 → 颜色（RGB 0~1）
GRADE_COLOR = {
    "可信": (0.09, 0.48, 0.35),
    "基本可信": (0.09, 0.48, 0.35),
    "需关注": (0.71, 0.33, 0.04),
    "高风险": (0.86, 0.15, 0.15),
    "不可评": (0.42, 0.45, 0.50),
}
_GRAY = (0.45, 0.48, 0.53)
_DARK = (0.10, 0.12, 0.16)
_AMBER = (0.71, 0.33, 0.04)
_RED = (0.86, 0.15, 0.15)


# ============================================================
#  报告模型
# ============================================================


class ReportSection(BaseModel):
    title: str
    lines: list[str] = Field(default_factory=list)


class ReportModel(BaseModel):
    title: str = "订单农业合同可信性评估报告"
    contract_title: str = ""
    doc_id: str = ""
    generated_at: str = ""
    verdict: OverallVerdict
    sections: list[ReportSection] = Field(default_factory=list)
    data_sources: list[str] = Field(default_factory=list)
    disclaimer: str = ""


# ============================================================
#  装配（从四模块原始输出抽取明细）
# ============================================================


def _contract_section(contract: Any) -> ReportSection:
    lines: list[str] = []
    if getattr(contract, "contract_no", ""):
        lines.append(f"合同编号：{contract.contract_no}")
    for party in getattr(contract, "parties", []) or []:
        code = getattr(party, "credit_code", "")
        suffix = f"（统一社会信用代码 {code}）" if code else ""
        lines.append(f"{getattr(party, 'role', '')}：{getattr(party, 'name', '')}{suffix}")
    for field in getattr(contract, "fields", []) or []:
        if getattr(field, "value", ""):
            lines.append(f"{field.label}：{field.value}")
    return ReportSection(title="合同要点", lines=lines)


def _standard_section(standard: dict | None) -> ReportSection | None:
    if not standard:
        return None
    req = standard.get("required_clauses") or {}
    comp = standard.get("compliance") or {}
    lines: list[str] = []
    if req.get("summary"):
        lines.append(req["summary"])
    if comp.get("summary"):
        lines.append(comp["summary"])
    for item in req.get("missing_required") or []:
        lines.append(f"缺失必备条款：{item.get('name', '')}（风险 {item.get('risk', '')}）")
    for item in req.get("missing_recommended") or []:
        lines.append(f"缺失建议条款：{item.get('name', '')}（风险 {item.get('risk', '')}）")
    for item in comp.get("violations") or []:
        lines.append(
            "偏离依据：{metric} 合同 {cv}{cu} vs {code} {clause} {sv}{su}（{basis}）".format(
                metric=item.get("metric", ""),
                cv=item.get("contract_value", ""),
                cu=item.get("contract_unit", ""),
                code=item.get("standard_code", ""),
                clause=item.get("clause_no", ""),
                sv=item.get("standard_value", ""),
                su=item.get("standard_unit", ""),
                basis=item.get("basis", "依据"),
            )
        )
    return ReportSection(title="标准校验明细", lines=lines)


def _parties_section(verify: Any) -> ReportSection | None:
    if verify is None:
        return None
    lines: list[str] = []
    if getattr(verify, "notice", ""):
        lines.append(f"核验说明：{verify.notice}")
    for party in getattr(verify, "parties", []) or []:
        who = f"{getattr(party, 'role', '')} {getattr(party, 'name', '')}".strip()
        lines.append(f"{who}：综合判定 {getattr(party, 'overall', '')}（工商查询 {getattr(party, 'l2_status', '')}）")
        for item in getattr(party, "l1", []) or []:
            lines.append(f"　· {item.label}：{item.status}——{item.detail}")
        if getattr(party, "l2_note", ""):
            lines.append(f"　· 工商信息：{party.l2_note}")
        for item in getattr(party, "l3", []) or []:
            lines.append(f"　· {item.label}：{item.status}——{item.detail}")
    return ReportSection(title="主体核验明细", lines=lines)


def _simulation_section(simulation: Any) -> ReportSection | None:
    if simulation is None:
        return None
    lines: list[str] = []
    if getattr(simulation, "notice", ""):
        lines.append(f"模拟说明：{simulation.notice}")
    for scenario in getattr(simulation, "scenarios", []) or []:
        label = getattr(scenario, "label", "") or getattr(scenario, "scenario", "")
        flag = "" if getattr(scenario, "computable", False) else "（无法量化）"
        lines.append(f"{label}{flag}：{getattr(scenario, 'summary', '')}")
    return ReportSection(title="履约模拟明细", lines=lines)


def _data_sources(standard: dict | None, verify: Any) -> list[str]:
    sources: list[str] = []
    if standard:
        index = standard.get("index") or {}
        embedding = (index.get("embedding") or {}).get("name")
        if embedding:
            sources.append(f"标准条文召回：{embedding}")
    provider = getattr(verify, "provider", "") if verify is not None else ""
    if provider:
        sources.append(f"工商数据来源：{provider}")
    return sources


def build_report(doc_id: str) -> ReportModel | None:
    """取数 → 归并 → 装配报告模型。"""
    collected = service.collect(doc_id)
    if collected is None:
        return None
    contract, standard, verify, simulation, errors = collected
    verdict = service.build_verdict(
        contract, standard, verify, simulation, errors=errors, doc_id=doc_id
    )

    sections: list[ReportSection] = []
    contract_section = _contract_section(contract)
    if contract_section.lines:
        sections.append(contract_section)
    for section in (
        _standard_section(standard),
        _parties_section(verify),
        _simulation_section(simulation),
    ):
        if section and section.lines:
            sections.append(section)

    return ReportModel(
        contract_title=verdict.contract_title,
        doc_id=doc_id,
        generated_at=verdict.generated_at,
        verdict=verdict,
        sections=sections,
        data_sources=_data_sources(standard, verify),
        disclaimer=verdict.disclaimer,
    )


# ============================================================
#  PDF 渲染
# ============================================================


def _wrap(text: str, size: float, max_width: float) -> list[str]:
    """按像素宽度做中英文混排换行（PyMuPDF 不自动折行）。"""
    out: list[str] = []
    for paragraph in (text or "").split("\n"):
        if not paragraph:
            out.append("")
            continue
        current = ""
        for ch in paragraph:
            if fitz.get_text_length(current + ch, fontname=FONT, fontsize=size) <= max_width:
                current += ch
            else:
                out.append(current)
                current = ch
        out.append(current)
    return out


class _PdfWriter:
    def __init__(self) -> None:
        self.doc = fitz.open()
        self.page = self.doc.new_page(width=A4_W, height=A4_H)
        self.y = MARGIN_Y

    def _room(self, height: float) -> None:
        if self.y + height > A4_H - MARGIN_Y:
            self.page = self.doc.new_page(width=A4_W, height=A4_H)
            self.y = MARGIN_Y

    def text(
        self,
        value: str,
        size: float = 10.5,
        color: tuple[float, float, float] = _DARK,
        indent: float = 0,
        gap: float = 5,
    ) -> None:
        max_width = A4_W - 2 * MARGIN_X - indent
        for line in _wrap(value, size, max_width):
            self._room(size + gap)
            self.page.insert_text(
                (MARGIN_X + indent, self.y), line, fontname=FONT, fontsize=size, color=color
            )
            self.y += size + gap
        self.y += 1

    def centered(
        self, value: str, size: float, color: tuple[float, float, float] = _DARK
    ) -> None:
        self._room(size + 8)
        width = fitz.get_text_length(value, fontname=FONT, fontsize=size)
        self.page.insert_text(
            ((A4_W - width) / 2, self.y), value, fontname=FONT, fontsize=size, color=color
        )
        self.y += size + 8

    def rule(self) -> None:
        self._room(10)
        self.page.draw_line(
            fitz.Point(MARGIN_X, self.y), fitz.Point(A4_W - MARGIN_X, self.y),
            color=(0.82, 0.84, 0.87), width=0.7,
        )
        self.y += 9

    def gap(self, height: float = 8) -> None:
        self.y += height

    def finish(self) -> bytes:
        data = self.doc.tobytes()
        self.doc.close()
        return data


def render_pdf(model: ReportModel) -> bytes:
    writer = _PdfWriter()
    verdict = model.verdict
    grade_color = GRADE_COLOR.get(verdict.grade, _GRAY)

    writer.centered(model.title, 20)
    writer.gap(6)
    writer.centered(f"总体可信度：{verdict.grade}", 15, grade_color)
    writer.gap(4)
    meta = f"{model.contract_title or '未命名合同'}    文档 ID {model.doc_id}"
    writer.text(meta, 9.5, _GRAY, gap=3)
    writer.text(f"生成时间：{model.generated_at}", 9.5, _GRAY, gap=3)
    writer.gap(4)
    writer.rule()

    writer.text("一、总体结论", 13, _DARK, gap=7)
    writer.text(verdict.grade_reason, 10.5)
    if verdict.narrative:
        writer.gap(2)
        writer.text(verdict.narrative, 10)

    if verdict.blockers:
        writer.gap(2)
        writer.text("否决项（客观硬伤）", 11, _RED, gap=6)
        for blocker in verdict.blockers:
            writer.text(f"· {blocker.label}", 10.5, _RED)
            if blocker.detail:
                writer.text(blocker.detail, 9.5, _GRAY, indent=14)
            if blocker.evidence:
                writer.text(f"依据：{blocker.evidence}", 9.5, _GRAY, indent=14)

    writer.gap(6)
    writer.text("二、分项判定", 13, _DARK, gap=7)
    for dim in verdict.dimensions:
        line = f"● {dim.label}：{dim.level}"
        if dim.score:
            line += f"（{dim.score}/100）"
        writer.text(line, 11, GRADE_COLOR.get(dim.level, _DARK))
        if dim.summary:
            writer.text(dim.summary, 9.5, _DARK, indent=14)
        for finding in dim.findings:
            writer.text(f"· {finding}", 9, _GRAY, indent=14)
        if dim.notice:
            writer.text(f"未完整评估：{dim.notice}", 9, _AMBER, indent=14)

    if model.sections:
        writer.gap(6)
        writer.text("三、模块明细", 13, _DARK, gap=7)
        for section in model.sections:
            writer.text(f"【{section.title}】", 11, _DARK, gap=5)
            for entry in section.lines:
                writer.text(entry, 9.5, _GRAY, indent=10)

    if verdict.notices:
        writer.gap(6)
        writer.text("四、未评估与提示", 13, _AMBER, gap=7)
        for note in verdict.notices:
            writer.text(f"· {note}", 9.5, _AMBER)

    if model.data_sources:
        writer.gap(6)
        writer.text("数据来源", 10, _GRAY, gap=5)
        for source in model.data_sources:
            writer.text(f"· {source}", 9, _GRAY, indent=10)

    writer.gap(10)
    writer.rule()
    writer.text(model.disclaimer, 8.5, _GRAY, gap=3)

    return writer.finish()


# ============================================================
#  Word 渲染
# ============================================================


def _set_cjk_font(run, name: str = "宋体") -> None:
    from docx.oxml.ns import qn

    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def render_docx(model: ReportModel) -> bytes:
    import io

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    verdict = model.verdict
    grade_rgb = RGBColor(*[int(c * 255) for c in GRADE_COLOR.get(verdict.grade, _GRAY)])

    def para(text: str, size: float = 10.5, bold: bool = False, color=None, align=None, indent=None):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        if indent is not None:
            p.paragraph_format.left_indent = Pt(indent)
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        _set_cjk_font(run)
        return p

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(model.title)
    run.bold = True
    run.font.size = Pt(20)
    _set_cjk_font(run)

    grade_line = doc.add_paragraph()
    grade_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = grade_line.add_run(f"总体可信度：{verdict.grade}")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = grade_rgb
    _set_cjk_font(run)

    para(f"{model.contract_title or '未命名合同'}    文档 ID {model.doc_id}", 9, color=RGBColor(0x64, 0x74, 0x8B), align=WD_ALIGN_PARAGRAPH.CENTER)
    para(f"生成时间：{model.generated_at}", 9, color=RGBColor(0x64, 0x74, 0x8B), align=WD_ALIGN_PARAGRAPH.CENTER)

    def heading(text: str, size: float = 13, color=None):
        para(text, size, bold=True, color=color)

    heading("一、总体结论")
    para(verdict.grade_reason)
    if verdict.narrative:
        para(verdict.narrative, 10)

    if verdict.blockers:
        heading("否决项（客观硬伤）", 11, RGBColor(0xDC, 0x26, 0x26))
        for blocker in verdict.blockers:
            para(f"· {blocker.label}", 10.5, bold=True, color=RGBColor(0xDC, 0x26, 0x26))
            if blocker.detail:
                para(blocker.detail, 9.5, indent=14)
            if blocker.evidence:
                para(f"依据：{blocker.evidence}", 9.5, indent=14)

    heading("二、分项判定")
    for dim in verdict.dimensions:
        color = RGBColor(*[int(c * 255) for c in GRADE_COLOR.get(dim.level, _DARK)])
        label = f"● {dim.label}：{dim.level}"
        if dim.score:
            label += f"（{dim.score}/100）"
        para(label, 11, bold=True, color=color)
        if dim.summary:
            para(dim.summary, 9.5, indent=14)
        for finding in dim.findings:
            para(f"· {finding}", 9, indent=14)
        if dim.notice:
            para(f"未完整评估：{dim.notice}", 9, color=RGBColor(0xB4, 0x53, 0x09), indent=14)

    if model.sections:
        heading("三、模块明细")
        for section in model.sections:
            para(f"【{section.title}】", 11, bold=True)
            for entry in section.lines:
                para(entry, 9.5, indent=10)

    if verdict.notices:
        heading("四、未评估与提示", color=RGBColor(0xB4, 0x53, 0x09))
        for note in verdict.notices:
            para(f"· {note}", 9.5, color=RGBColor(0xB4, 0x53, 0x09))

    if model.data_sources:
        heading("数据来源", 10)
        for source in model.data_sources:
            para(f"· {source}", 9, indent=10)

    para(model.disclaimer, 8.5, color=RGBColor(0x64, 0x74, 0x8B))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ============================================================
#  导出到文件（供 MCP / 脚本用）
# ============================================================


def export_to_file(doc_id: str, fmt: str = "pdf") -> tuple[Path, int] | None:
    """生成报告文件到 data/reports/，返回 (路径, 字节数)。合同不存在返回 None。"""
    model = build_report(doc_id)
    if model is None:
        return None
    if fmt == "docx":
        data, ext = render_docx(model), "docx"
    elif fmt == "pdf":
        data, ext = render_pdf(model), "pdf"
    else:
        raise ValueError(f"不支持的报告格式：{fmt}（仅 pdf / docx）")

    directory = settings.data_dir / "reports"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"report_{doc_id}.{ext}"
    path.write_bytes(data)
    return path, len(data)
